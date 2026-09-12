# app/ai/document_report.py
# Renders a Controlled Document Maturity Assessment result into an audit-grade
# .docx: real Word-native multilevel heading numbering (not manually-typed),
# a real TOC field, header/footer with real page-number fields, Aptos font,
# autofit-to-window tables with repeating header rows, landscape sections for
# the wide tables, and heat-map colored scores. Sibling to report_generator.py,
# which handles the unrelated Industry/Standard/Template questionnaire reports.

import re
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.ai.report_generator import REPORTS_DIR

NAVY = RGBColor(0x0F, 0x3A, 0x7D)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1F, 0x29, 0x37)

# Score -> (fill hex, meaning, text color) per the rubric's heat-map scale.
HEAT_MAP = {
    5: ("1F4E79", "Excellent", WHITE),
    4: ("2E7D32", "Strong", WHITE),
    3: ("F2C94C", "Adequate", DARK),
    2: ("F2994A", "Weak", DARK),
    1: ("EB5757", "Critical Deficiency", WHITE),
}

PRIORITY_ORDER = ["Critical", "High", "Medium", "Low"]


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------

def _hex_rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _add_multilevel_numbering(doc: Document) -> str:
    """Defines a real Word multilevel numbered list (1 / 1.1 / 1.1.1) and
    returns its numId, so heading styles can be linked to it — this is what
    makes the numbering automatic/native instead of manually-typed text."""
    numbering_part = doc.part.numbering_part
    numbering_elem = numbering_part.element

    abstract_num_id = "100"
    num_id = "100"

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), abstract_num_id)
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "hybridMultilevel")
    abstract_num.append(multilevel)

    level_texts = ["%1", "%1.%2", "%1.%2.%3"]
    for i, level_text in enumerate(level_texts):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(i))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), level_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(360 * (i + 1)))
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        r_pr = OxmlElement("w:rPr")
        r_pr.append(OxmlElement("w:b"))
        for el in (start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr):
            lvl.append(el)
        abstract_num.append(lvl)

    numbering_elem.append(abstract_num)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    abstract_num_id_ref = OxmlElement("w:abstractNumId")
    abstract_num_id_ref.set(qn("w:val"), abstract_num_id)
    num.append(abstract_num_id_ref)
    numbering_elem.append(num)

    return num_id


def _insert_numpr(p_pr, ilvl: int, num_id: str):
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    insert_at = 0
    leading_tags = {"pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl"}
    for i, child in enumerate(p_pr):
        tag = child.tag.split("}")[-1]
        if tag in leading_tags:
            insert_at = i + 1
        else:
            break
    p_pr.insert(insert_at, num_pr)
    return num_pr


def _link_heading_numbering(doc: Document, num_id: str):
    for level, style_name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        style = doc.styles[style_name]
        p_pr = style.element.get_or_add_pPr()
        _insert_numpr(p_pr, level, num_id)


def _suppress_numbering(paragraph):
    """Overrides a specific paragraph to numId=0 (Word's 'no numbering'
    convention) — used for the Table of Contents heading, which shouldn't
    be part of the numbered outline even though it uses the Heading 1 style."""
    p_pr = paragraph._p.get_or_add_pPr()
    _insert_numpr(p_pr, 0, "0")


def _configure_document_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.2

    for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(size + 4)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True  # prevents orphaned headings at page bottom


def _heading(doc: Document, level: int, text: str):
    style = f"Heading {level}"
    return doc.add_paragraph(text, style=style)


def _insert_toc_field(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click here and choose Update Field to generate the table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r.append(fld_begin)
    r.append(instr)
    r.append(fld_separate)
    r.append(placeholder)
    r.append(fld_end)


def _add_field_run(paragraph, instr_text: str):
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_end)
    return run


def _add_header_footer(doc: Document, document_title: str):
    section = doc.sections[0]

    header_p = section.header.paragraphs[0]
    header_run = header_p.add_run(f"MaturityAssess | {document_title}")
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = GREY

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_str = datetime.now(timezone.utc).strftime("%B %d, %Y")
    lead_run = footer_p.add_run(f"Confidential | {date_str} | Page ")
    lead_run.font.size = Pt(9)
    lead_run.font.color.rgb = GREY
    page_run = _add_field_run(footer_p, "PAGE")
    page_run.font.size = Pt(9)
    page_run.font.color.rgb = GREY
    mid_run = footer_p.add_run(" of ")
    mid_run.font.size = Pt(9)
    mid_run.font.color.rgb = GREY
    total_run = _add_field_run(footer_p, "NUMPAGES")
    total_run.font.size = Pt(9)
    total_run.font.color.rgb = GREY


def _set_table_autofit_window(table):
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    tbl_pr.append(tbl_w)
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "autofit")
    tbl_pr.append(tbl_layout)


def _repeat_header_row(table):
    """Marks the header row to repeat on continuation pages if the table
    breaks across a page boundary."""
    tr = table.rows[0]._tr
    tr_pr = tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _shade_cell(cell, hex_color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, color: RGBColor | None = None, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    if color is not None:
        run.font.color.rgb = color
    run.font.bold = bold


def _make_table(doc: Document, headers: list[str], style: str = "Light Grid Accent 1"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    _set_table_autofit_window(table)
    _repeat_header_row(table)
    for i, label in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], label, bold=True)
    return table


def _start_landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    return section


def _start_portrait_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    return section


def _bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _kv_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    _set_table_autofit_window(table)
    for label, value in rows:
        row = table.add_row().cells
        _set_cell_text(row[0], label, bold=True)
        _set_cell_text(row[1], value or "—")
    return table


def _score_line(doc: Document, score: int, prefix: str = "Score"):
    fill, meaning, _ = HEAT_MAP.get(score, ("6B7280", "Unscored", DARK))
    p = doc.add_paragraph()
    run = p.add_run(f"{prefix}: {score}/5 — {meaning}")
    run.bold = True
    run.font.color.rgb = _hex_rgb(fill)
    return p


# ---------------------------------------------------------------------------
# Section builders
# ---------------------------------------------------------------------------

def _cover_and_toc(doc: Document, document_title: str, context: dict, original_filename: str):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_title)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = NAVY

    badge = doc.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge_run = badge.add_run("CONTROLLED DOCUMENT MATURITY ASSESSMENT")
    badge_run.bold = True
    badge_run.font.size = Pt(11)
    badge_run.font.color.rgb = RGBColor(0x00, 0xA8, 0xE8)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_lines = [
        context.get("organization_name") or "",
        context.get("process_name") or "",
        f"Source document: {original_filename}",
        f"Assessment date: {datetime.now(timezone.utc).strftime('%B %d, %Y')}",
        "Prepared by MaturityAssess",
    ]
    for line in [m for m in meta_lines if m]:
        r = meta_p.add_run(line + "\n")
        r.font.color.rgb = GREY

    doc.add_page_break()

    toc_heading = doc.add_paragraph("Table of Contents", style="Heading 1")
    _suppress_numbering(toc_heading)
    _insert_toc_field(doc)
    doc.add_page_break()


def _section_executive_summary(doc: Document, result: dict):
    _heading(doc, 1, "Executive Summary")

    _heading(doc, 2, "Assessment Overview")
    doc.add_paragraph(
        "This report presents a full evidence-based maturity assessment of the "
        "submitted document against a fixed 15-category audit rubric."
    )

    _heading(doc, 2, "Overall Maturity Result")
    _kv_table(doc, [
        ("Overall Maturity Score", f"{result.get('overall_maturity_score', 0)}/5"),
        ("Overall Maturity Level", result.get("overall_maturity_level", "")),
        ("Overall Readiness", f"{result.get('overall_readiness_pct', 0)}%"),
    ])

    _heading(doc, 2, "Compliance Readiness")
    doc.add_paragraph(f"{result.get('compliance_readiness_pct', 0)}% — see Section 4 for the full compliance profile.")

    _heading(doc, 2, "Evidence Confidence")
    doc.add_paragraph(
        f"Evidence coverage: {result.get('evidence_coverage_pct', 0)}%. "
        f"Overall assessment confidence: {result.get('assessment_confidence', '')}."
    )

    _heading(doc, 2, "Key Strengths")
    _bullets(doc, result.get("key_strengths", []))

    _heading(doc, 2, "Critical Gaps")
    critical_gaps = [g for g in result.get("gap_analysis", []) if g.get("priority") in ("Critical", "High")]
    _bullets(doc, [f"{g.get('category', '')}: {g.get('requirement', '')}" for g in critical_gaps] or ["None identified."])

    _heading(doc, 2, "Priority Recommendations")
    _bullets(doc, [r.get("recommendation", "") for r in result.get("top_recommendations", [])[:5]])

    _heading(doc, 2, "Management Conclusion")
    doc.add_paragraph(
        f"Current maturity: {result.get('current_maturity_level', '')}. "
        f"Full management conclusion in Section 12."
    )


def _section_scope_methodology(doc: Document, document_title: str, original_filename: str):
    _heading(doc, 1, "Assessment Scope and Methodology")

    _heading(doc, 2, "Objective")
    doc.add_paragraph(
        "Evaluate the maturity and effectiveness of the submitted document as both a "
        "controlled management-system document and an operational business tool."
    )

    _heading(doc, 2, "Source Document")
    doc.add_paragraph(f"{document_title} ({original_filename}).")

    _heading(doc, 2, "Scope")
    doc.add_paragraph(
        "The assessment covers the content of the submitted document only, evaluated "
        "against the 15 fixed categories in Section 6."
    )

    _heading(doc, 2, "Method")
    doc.add_paragraph(
        "Context (organization, industry, process, governance, and compliance profile) "
        "was extracted from the document and reviewed/corrected by the user before "
        "scoring. Scoring was then performed against that reviewed context and the "
        "document's actual text."
    )

    _heading(doc, 2, "Maturity Model")
    _make_table(doc, ["Level", "Description"])
    table = doc.tables[-1]
    for level, desc in (
        (1, "Initial / Ad Hoc"), (2, "Repeatable"), (3, "Defined"),
        (4, "Managed"), (5, "Optimized / Best Practice"),
    ):
        row = table.add_row().cells
        _set_cell_text(row[0], str(level))
        _set_cell_text(row[1], desc)

    _heading(doc, 2, "Evidence Methodology")
    doc.add_paragraph(
        "Every finding is classified as Explicit Evidence (directly stated), Strong "
        "Indication (reasonably supported but not explicit), Not Evidenced "
        "(insufficient basis in the document), or External Benchmark (derived from an "
        "applicable external standard, not from the document itself). Mentioning a "
        "regulation or standard is never treated as proof of compliance with it."
    )

    _heading(doc, 2, "Scoring Method")
    doc.add_paragraph(
        "Each category is scored 1-5 against the maturity model above, with a stated "
        "confidence level (High/Medium/Low) reflecting the quantity and quality of "
        "evidence available, not just whether the topic was mentioned."
    )

    _heading(doc, 2, "Limitations")
    doc.add_paragraph(
        "This assessment is based solely on the content of the submitted document. "
        "Source citations for PDF pages are page-based; DOCX citations are "
        "section-based, since Word documents have no fixed page layout before "
        "rendering. OCR-derived text (where used, for scanned/image-only PDF pages) "
        "may contain transcription errors. A finding of 'Not Evidenced' means the "
        "document does not demonstrate the practice — not that the practice doesn't "
        "exist in the organization."
    )


def _section_organization_profile(doc: Document, context: dict):
    _heading(doc, 1, "Organization and Process Profile")

    _heading(doc, 2, "Organization")
    _kv_table(doc, [
        ("Organization", context.get("organization_name", "")), ("Business Unit", context.get("business_unit", "")),
        ("Department", context.get("department", "")), ("Function", context.get("function", "")),
        ("Location", context.get("location", "")),
    ])

    _heading(doc, 2, "Industry")
    _kv_table(doc, [
        ("Industry", context.get("industry", "")), ("Sub-Industry", context.get("sub_industry", "")),
        ("Business Model", context.get("business_model", "")),
        ("Operating Environment", context.get("operating_environment", "")),
    ])

    _heading(doc, 2, "Business Function")
    doc.add_paragraph(context.get("function") or "Not evidenced in the source document.")

    _heading(doc, 2, "Process")
    _kv_table(doc, [
        ("Name", context.get("process_name", "")), ("Purpose", context.get("process_purpose", "")),
        ("Scope", context.get("process_scope", "")), ("Boundaries", context.get("process_boundaries", "")),
        ("Inputs", context.get("process_inputs", "")), ("Outputs", context.get("process_outputs", "")),
        ("Activities", context.get("process_activities", "")), ("Process Owner", context.get("process_owner", "")),
    ])

    _heading(doc, 2, "Geographic Scope")
    doc.add_paragraph(context.get("location") or "Not evidenced in the source document.")

    _heading(doc, 2, "Stakeholders")
    doc.add_paragraph(context.get("governance_roles") or "Not evidenced in the source document.")

    _heading(doc, 2, "Systems and Technology")
    systems = context.get("systems", []) or []
    _bullets(doc, systems or ["Not evidenced in the source document."])
    if context.get("automation_notes"):
        doc.add_paragraph(context["automation_notes"])

    _heading(doc, 2, "Process Interfaces")
    doc.add_paragraph(
        "See Section 6.12 (Process Integration) for how this process integrates with "
        "related procedures, systems, and enterprise processes."
    )


def _section_compliance_profile(doc: Document, context: dict):
    _heading(doc, 1, "Regulatory, Compliance and Standards Profile")
    items = context.get("compliance_profile", []) or []

    for n, label, keyword in (
        (1, "Regulations", "regulat"), (2, "Standards", "standard"), (3, "Certifications", "certif"),
    ):
        _heading(doc, 2, label)
        matches = [i for i in items if keyword in (i.get("req_type", "") or "").lower()]
        if matches:
            _bullets(doc, [f"{i.get('requirement', '')} — {i.get('status', '')}" for i in matches])
        else:
            doc.add_paragraph("Not evidenced in the source document.")

    _heading(doc, 2, "Internal Requirements")
    internal = [i for i in items if "internal" in (i.get("req_type", "") or "").lower() or "polic" in (i.get("req_type", "") or "").lower()]
    _bullets(doc, [f"{i.get('requirement', '')} — {i.get('status', '')}" for i in internal] or ["Not evidenced in the source document."])

    _heading(doc, 2, "Applicability")
    doc.add_paragraph("Full requirement-by-requirement applicability, evidence, and status:")
    if items:
        table = _make_table(doc, ["Requirement", "Type", "Applicability", "Evidence", "Status", "Source"])
        for i in items:
            row = table.add_row().cells
            _set_cell_text(row[0], i.get("requirement", ""))
            _set_cell_text(row[1], i.get("req_type", ""))
            _set_cell_text(row[2], i.get("applicability", ""))
            _set_cell_text(row[3], i.get("evidence", ""))
            _set_cell_text(row[4], i.get("status", ""))
            _set_cell_text(row[5], i.get("source", ""))
    else:
        doc.add_paragraph("No regulatory, standard, or compliance references were found in the source document.")

    _heading(doc, 2, "Evidence")
    doc.add_paragraph("Evidence for each requirement is captured in the table above, cited to its source where identifiable.")

    _heading(doc, 2, "Compliance Gaps")
    gaps = [i for i in items if i.get("status") in ("Not Evidenced", "Gap Identified", "Partially Addressed")]
    _bullets(doc, [f"{i.get('requirement', '')} — {i.get('status', '')}" for i in gaps] or ["None identified."])


def _section_maturity_results(doc: Document, result: dict):
    _heading(doc, 1, "Maturity Assessment Results")

    _heading(doc, 2, "Overall Score")
    doc.add_paragraph(f"{result.get('overall_maturity_score', 0)}/5")

    _heading(doc, 2, "Overall Level")
    doc.add_paragraph(result.get("overall_maturity_level", ""))

    _heading(doc, 2, "Category Summary")
    table = _make_table(doc, ["Category", "Score", "Confidence"])
    for c in result.get("category_scores", []):
        row = table.add_row().cells
        _set_cell_text(row[0], c.get("category", ""))
        _set_cell_text(row[1], str(c.get("score", "")))
        _set_cell_text(row[2], c.get("confidence", ""))

    _heading(doc, 2, "Heat Map")
    legend = _make_table(doc, [f"{s} — {HEAT_MAP[s][1]}" for s in (5, 4, 3, 2, 1)], style="Table Grid")
    for i, score in enumerate((5, 4, 3, 2, 1)):
        fill, _, color = HEAT_MAP[score]
        cell = legend.rows[0].cells[i]
        _shade_cell(cell, fill)
        cell.paragraphs[0].runs[0].font.color.rgb = color

    heat_table = _make_table(doc, ["Category", "Score", "Evidence Found", "Gaps Identified", "Recommendations", "Heat Map"])
    for c in result.get("category_scores", []):
        row = heat_table.add_row().cells
        score = int(c.get("score", 0) or 0)
        fill, meaning, color = HEAT_MAP.get(score, ("FFFFFF", "Unscored", DARK))
        _set_cell_text(row[0], c.get("category", ""))
        _set_cell_text(row[1], str(score), color=color, bold=True)
        _shade_cell(row[1], fill)
        _set_cell_text(row[2], c.get("evidence_found", ""))
        _set_cell_text(row[3], c.get("gaps_identified", ""))
        _set_cell_text(row[4], c.get("recommendation", ""))
        _set_cell_text(row[5], meaning, color=color, bold=True)
        _shade_cell(row[5], fill)

    _heading(doc, 2, "Evidence Coverage")
    doc.add_paragraph(f"{result.get('evidence_coverage_pct', 0)}%")

    _heading(doc, 2, "Confidence")
    doc.add_paragraph(result.get("assessment_confidence", ""))


def _section_detailed_assessment(doc: Document, result: dict, categories: list[dict]):
    _heading(doc, 1, "Detailed Maturity Assessment")
    assess_by_name = {c["name"]: c["assess"] for c in categories}
    for c in result.get("category_scores", []):
        _heading(doc, 2, c.get("category", ""))
        doc.add_paragraph(assess_by_name.get(c.get("category", ""), "")).runs[0].italic = True
        score = int(c.get("score", 0) or 0)
        _score_line(doc, score)
        doc.add_paragraph(f"Confidence: {c.get('confidence', '')}    Source: {c.get('source', '')}")
        _labeled = doc.add_paragraph()
        _labeled.add_run("Evidence: ").bold = True
        _labeled.add_run(c.get("evidence_found", ""))
        gaps_p = doc.add_paragraph()
        gaps_p.add_run("Gaps: ").bold = True
        gaps_p.add_run(c.get("gaps_identified", ""))
        if c.get("strengths_and_risks"):
            sr_p = doc.add_paragraph()
            sr_p.add_run("Strengths & Risks: ").bold = True
            sr_p.add_run(c["strengths_and_risks"])
        rec_p = doc.add_paragraph()
        rec_p.add_run("Recommendation: ").bold = True
        rec_p.add_run(c.get("recommendation", ""))


def _section_gap_analysis(doc: Document, result: dict):
    _heading(doc, 1, "Detailed Gap Analysis")
    gaps = result.get("gap_analysis", [])
    for priority in PRIORITY_ORDER:
        _heading(doc, 2, priority)
        matches = [g for g in gaps if g.get("priority") == priority]
        if not matches:
            doc.add_paragraph("None identified.")
            continue
        table = _make_table(doc, ["Gap ID", "Category", "Requirement", "Gap Detail", "Impact"])
        for g in matches:
            row = table.add_row().cells
            _set_cell_text(row[0], g.get("gap_id", ""))
            _set_cell_text(row[1], g.get("category", ""))
            _set_cell_text(row[2], g.get("requirement", ""))
            _set_cell_text(row[3], f"{g.get('gap_detail', '')} (Source: {g.get('source', '')})")
            _set_cell_text(row[4], g.get("impact", ""))


def _section_recommendations(doc: Document, result: dict):
    _heading(doc, 1, "Top Improvement Recommendations")
    for i, r in enumerate(result.get("top_recommendations", []), start=1):
        _heading(doc, 2, f"Recommendation {i}" + ("  [QUICK WIN]" if r.get("quick_win") else ""))
        doc.add_paragraph(r.get("recommendation", ""))
        _kv_table(doc, [
            ("Business Benefit", r.get("business_benefit", "")),
            ("Risk if Not Addressed", r.get("risk_if_not_addressed", "")),
            ("Priority", r.get("priority", "")), ("Estimated Maturity Gain", r.get("estimated_maturity_gain", "")),
            ("Evidence Basis", r.get("evidence_basis", "")),
        ])


def _section_quick_wins(doc: Document, result: dict):
    _heading(doc, 1, "Quick Wins")
    wins = [r for r in result.get("top_recommendations", []) if r.get("quick_win")]
    if not wins:
        doc.add_paragraph("No recommendations were flagged as quick wins in this assessment.")
        return
    for r in wins:
        p = doc.add_paragraph(style="List Bullet")
        badge = p.add_run("QUICK WIN: ")
        badge.bold = True
        badge.font.color.rgb = _hex_rgb(HEAT_MAP[4][0])
        p.add_run(r.get("recommendation", ""))


def _section_roadmap(doc: Document, result: dict):
    _heading(doc, 1, "30/60/90-Day Roadmap")
    for label, key in (
        ("0-30 Days", "roadmap_day_30"), ("31-60 Days", "roadmap_day_60"),
        ("61-90 Days", "roadmap_day_90"), ("Beyond 90 Days", "roadmap_beyond_90"),
    ):
        _heading(doc, 2, label)
        items = result.get(key, [])
        if not items:
            doc.add_paragraph("No actions identified for this horizon.")
            continue
        _bullets(doc, items)


def _section_target_state(doc: Document, result: dict):
    _heading(doc, 1, "Target Maturity State")
    doc.add_paragraph(result.get("target_maturity_level", ""))


def _section_management_conclusion(doc: Document, result: dict):
    _heading(doc, 1, "Management Conclusion")

    _heading(doc, 2, "Maturity Summary")
    _kv_table(doc, [
        ("Current Maturity Level", result.get("current_maturity_level", "")),
        ("Target Maturity Level", "Level 5 — Optimized / Best Practice"),
        ("Overall Maturity Score", f"{result.get('overall_maturity_score', 0)}/5"),
    ])

    _heading(doc, 2, "Readiness Summary")
    _kv_table(doc, [
        ("Compliance Readiness", f"{result.get('compliance_readiness_pct', 0)}%"),
        ("Evidence Coverage", f"{result.get('evidence_coverage_pct', 0)}%"),
        ("Assessment Confidence", result.get("assessment_confidence", "")),
        ("Overall Readiness", f"{result.get('overall_readiness_pct', 0)}%"),
    ])

    _heading(doc, 2, "Critical Risks")
    critical = [g for g in result.get("gap_analysis", []) if g.get("priority") == "Critical"]
    _bullets(doc, [g.get("risk", "") for g in critical] or ["None identified."])

    _heading(doc, 2, "Priority Actions")
    _bullets(doc, [r.get("recommendation", "") for r in result.get("top_recommendations", [])[:5]])

    _heading(doc, 2, "Roadmap to Level 5")
    doc.add_paragraph("See Section 10 for the full 30/60/90/Beyond-90 roadmap, and Section 11 for the target state.")


def _section_appendices(doc: Document, result: dict, context: dict):
    _heading(doc, 1, "Appendices")

    _heading(doc, 2, "Evidence Register")
    table = _make_table(doc, ["Category", "Score", "Confidence", "Source", "Evidence Found"])
    for c in result.get("category_scores", []):
        row = table.add_row().cells
        _set_cell_text(row[0], c.get("category", ""))
        _set_cell_text(row[1], str(c.get("score", "")))
        _set_cell_text(row[2], c.get("confidence", ""))
        _set_cell_text(row[3], c.get("source", ""))
        _set_cell_text(row[4], c.get("evidence_found", ""))

    _heading(doc, 2, "Compliance Register")
    items = context.get("compliance_profile", []) or []
    if items:
        table = _make_table(doc, ["Requirement", "Status", "Source"])
        for i in items:
            row = table.add_row().cells
            _set_cell_text(row[0], i.get("requirement", ""))
            _set_cell_text(row[1], i.get("status", ""))
            _set_cell_text(row[2], i.get("source", ""))
    else:
        doc.add_paragraph("No compliance items were identified in the source document.")

    _heading(doc, 2, "Assessment Scorecard")
    table = _make_table(doc, ["Category", "Score"])
    for c in result.get("category_scores", []):
        row = table.add_row().cells
        _set_cell_text(row[0], c.get("category", ""))
        _set_cell_text(row[1], str(c.get("score", "")))

    _heading(doc, 2, "Recommendation Register")
    table = _make_table(doc, ["#", "Recommendation", "Priority", "Quick Win"])
    for i, r in enumerate(result.get("top_recommendations", []), start=1):
        row = table.add_row().cells
        _set_cell_text(row[0], str(i))
        _set_cell_text(row[1], r.get("recommendation", ""))
        _set_cell_text(row[2], r.get("priority", ""))
        _set_cell_text(row[3], "Yes" if r.get("quick_win") else "No")

    _heading(doc, 2, "Assumptions and Limitations")
    doc.add_paragraph(
        "This assessment is based solely on the content of the uploaded document, as "
        "reviewed and corrected by the user before scoring. DOCX source citations are "
        "section-based rather than page-based. OCR-derived text (where used) may "
        "contain transcription errors. A finding of 'Not Evidenced' means the "
        "document does not demonstrate the practice, not that the organization lacks "
        "it in practice."
    )

    _heading(doc, 2, "Source References")
    sources = set()
    for c in result.get("category_scores", []):
        if c.get("source"):
            sources.add(c["source"])
    for g in result.get("gap_analysis", []):
        if g.get("source"):
            sources.add(g["source"])
    for i in items:
        if i.get("source"):
            sources.add(i["source"])
    _bullets(doc, sorted(sources) or ["No source citations were captured for this assessment."])


# ---------------------------------------------------------------------------
# Top-level render
# ---------------------------------------------------------------------------

def render_document_assessment_docx(
    *, assessment_id: str, document_title: str, original_filename: str, result: dict, context: dict,
) -> str:
    from app.ai.document_maturity import DOCUMENT_MATURITY_CATEGORIES

    doc = Document()
    _configure_document_styles(doc)
    num_id = _add_multilevel_numbering(doc)
    _link_heading_numbering(doc, num_id)
    _add_header_footer(doc, document_title)

    _cover_and_toc(doc, document_title, context, original_filename)

    _section_executive_summary(doc, result)
    _section_scope_methodology(doc, document_title, original_filename)
    _section_organization_profile(doc, context)

    _start_landscape_section(doc)
    _section_compliance_profile(doc, context)
    _section_maturity_results(doc, result)

    _start_portrait_section(doc)
    _section_detailed_assessment(doc, result, DOCUMENT_MATURITY_CATEGORIES)

    _start_landscape_section(doc)
    _section_gap_analysis(doc, result)

    _start_portrait_section(doc)
    _section_recommendations(doc, result)
    _section_quick_wins(doc, result)
    _section_roadmap(doc, result)
    _section_target_state(doc, result)
    _section_management_conclusion(doc, result)
    _section_appendices(doc, result, context)

    file_path = REPORTS_DIR / f"{assessment_id}_document_assessment.docx"
    doc.save(str(file_path))
    return str(file_path)


def download_filename(document_title: str, created_at: datetime) -> str:
    safe = re.sub(r"[^A-Za-z0-9_\-]+", "_", document_title.strip())
    safe = re.sub(r"_+", "_", safe).strip("_") or "Document"
    return f"{safe}_{created_at.strftime('%y%m%d')}.docx"
