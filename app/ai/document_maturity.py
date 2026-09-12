# app/ai/document_maturity.py
# Controlled Document Maturity Assessment: scores an uploaded procedure, policy,
# or similar governance document against a fixed 15-category rubric. Standalone
# from the Industry/Standard/Template questionnaire flow in the rest of app/ai —
# there are no questions here, the document itself is the input.

import io

from pypdf import PdfReader
from docx import Document as DocxDocument

from app.ai.claude_client import generate_structured
from app.ai.prompt_templates import STYLE_DIRECTIVE
from app.schemas import GeneratedDocumentAssessment

MAX_DOCUMENT_CHARS = 100_000

# The 15 fixed evaluation categories, each with the assessment bullets from the
# source rubric. Fixed and ordered so every assessment covers exactly these
# categories, in this order, rather than leaving category selection to Claude.
DOCUMENT_MATURITY_CATEGORIES: list[dict] = [
    {
        "name": "Strategic Alignment",
        "assess": (
            "Alignment with organizational strategy, business objectives, corporate "
            "policies, governance requirements, and regulatory obligations. Whether "
            "clear links are established, purpose is explained, and business value "
            "is evident."
        ),
    },
    {
        "name": "Purpose and Scope",
        "assess": (
            "Clarity of purpose, scope completeness, boundaries defined, inclusions "
            "and exclusions defined, applicability specified."
        ),
    },
    {
        "name": "Process Definition",
        "assess": (
            "Whether the process is described logically, workflow clarity, inputs "
            "identified, outputs identified, activities described, sequence "
            "understandable — whether a new employee could execute the process "
            "successfully using only this document."
        ),
    },
    {
        "name": "Roles and Responsibilities",
        "assess": (
            "Ownership defined, accountability defined, authorities defined, "
            "escalation paths defined, delegation rules defined, whether RACI "
            "principles are evident, and whether ambiguity exists."
        ),
    },
    {
        "name": "Controls and Governance",
        "assess": (
            "Approval requirements, review requirements, monitoring activities, "
            "compliance requirements, management oversight, and the overall "
            "robustness of governance controls."
        ),
    },
    {
        "name": "Risk Integration",
        "assess": (
            "Whether the document addresses operational, compliance, safety, "
            "knowledge-loss, cybersecurity, and human-performance risks, and "
            "whether risk controls, mitigation measures, and escalation criteria "
            "are present."
        ),
    },
    {
        "name": "Document Control Maturity",
        "assess": (
            "Revision management, version control, approval records, review cycle, "
            "change control process, obsolete document controls, and alignment with "
            "documented-information requirements and document lifecycle controls."
        ),
    },
    {
        "name": "Process Performance Management",
        "assess": (
            "Presence of KPIs, targets, service levels, effectiveness measures, "
            "efficiency measures, and whether process success can be objectively "
            "measured."
        ),
    },
    {
        "name": "Knowledge Management",
        "assess": (
            "Lessons-learned integration, knowledge capture and sharing mechanisms, "
            "retention of critical knowledge, references to repositories, "
            "protection of knowledge, and vulnerability to knowledge loss."
        ),
    },
    {
        "name": "Human Factors and Usability",
        "assess": (
            "Readability, simplicity, logical structure, visual aids, process maps, "
            "flowcharts, tables, user friendliness, and whether the document "
            "supports human performance."
        ),
    },
    {
        "name": "Compliance Readiness",
        "assess": (
            "Regulatory alignment, standard alignment, internal requirements "
            "alignment, audit trail capability, evidence requirements, and audit "
            "readiness."
        ),
    },
    {
        "name": "Process Integration",
        "assess": (
            "How well the document integrates with related procedures, policies, "
            "manuals, systems, databases, and enterprise processes, and the degree "
            "of silo risk."
        ),
    },
    {
        "name": "Continuous Improvement Capability",
        "assess": (
            "Review triggers, feedback mechanisms, corrective actions, preventive "
            "actions, improvement opportunities, and whether continual improvement "
            "is embedded."
        ),
    },
    {
        "name": "Digital Readiness",
        "assess": (
            "Automation opportunities, workflow integration, system integration, AI "
            "enablement, data capture capability, and whether the process is future "
            "ready."
        ),
    },
    {
        "name": "Benchmark Against Industry Best Practice",
        "assess": (
            "Comparison against ISO management system practices, industry "
            "expectations, high reliability organizations, and world-class "
            "governance frameworks, and the resulting maturity gap."
        ),
    },
]


def extract_document_text(contents: bytes, filename: str) -> str:
    """Extracts plain text from an uploaded .pdf/.docx/.txt file. Raises
    ValueError on an unsupported or unreadable file."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "txt":
        text = contents.decode("utf-8", errors="ignore")
    elif ext == "docx":
        doc = DocxDocument(io.BytesIO(contents))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts)
    elif ext == "pdf":
        reader = PdfReader(io.BytesIO(contents))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    text = text.strip()
    if not text:
        raise ValueError(
            "No readable text could be extracted from this file. If it's a scanned "
            "or image-only PDF, this tool can't read it — text extraction only, no OCR."
        )

    if len(text) > MAX_DOCUMENT_CHARS:
        text = text[:MAX_DOCUMENT_CHARS] + "\n\n[Document truncated for length.]"

    return text


def document_assessment_prompt(document_title: str, document_text: str) -> tuple[str, str]:
    category_list = "\n".join(
        f"{i + 1}. {c['name']} — assess: {c['assess']}"
        for i, c in enumerate(DOCUMENT_MATURITY_CATEGORIES)
    )

    system = (
        "Act as a Senior Management Systems Auditor, Governance Assessor, Process "
        "Excellence Consultant, ISO Lead Auditor, and Document Control Expert, with "
        "expertise in ISO 9001, ISO 30401, ISO 55001, ISO 37301, ISO 37000, APQC "
        "Process Management, BPMM (Business Process Maturity Model), Document "
        "Governance, Enterprise Risk Management, Asset Intensive Industry Management "
        "Systems, and High Reliability Organization (HRO) Practices.\n\n"
        "Your task is to perform a comprehensive maturity assessment of a controlled "
        "document (a procedure, process description, management system manual, "
        "governance framework, standard, work instruction, or guideline) provided by "
        "the user. Evaluate the document as both a controlled management system "
        "document and an operational business tool: determine whether it merely "
        "exists, whether it is implemented effectively, and whether it demonstrates "
        "characteristics of a world-class process governance document.\n\n"
        "Score every one of the 15 fixed categories below on this maturity scale: "
        "1 = Initial/Ad Hoc, 2 = Repeatable, 3 = Defined, 4 = Managed, "
        "5 = Optimized/Best Practice. Ground every score, evidence_found, "
        "gaps_identified, and recommendation in what the document text actually "
        "says (or conspicuously fails to say) — never invent content, and never "
        "default every category to the same middle score. "
        + STYLE_DIRECTIVE
    )

    user = (
        f"Document title: {document_title}\n\n"
        f"Evaluate the document against exactly these 15 categories, in this order:\n"
        f"{category_list}\n\n"
        f"This is a full audit-grade assessment, not a summary — write it with the "
        f"depth of a consulting deliverable a client is paying for, not a checklist. "
        f"For each category, provide:\n"
        f"- score (1-5)\n"
        f"- evidence_found: a substantive paragraph (4-6 sentences) citing specific "
        f"content, wording, or structure actually in the document that supports this "
        f"score — quote or closely paraphrase where useful. If there's genuinely "
        f"little to cite, say so explicitly rather than padding.\n"
        f"- gaps_identified: a substantive paragraph (4-6 sentences) on what's "
        f"missing, weak, or ambiguous, specific to this document, with concrete "
        f"detail on why it matters (not just 'lacks X').\n"
        f"- recommendation: a detailed, actionable paragraph (3-5 sentences) — not a "
        f"one-line suggestion. Say specifically what to add or change, and how.\n\n"
        f"Also provide:\n"
        f"- overall_maturity_score (average across categories) and "
        f"overall_maturity_level\n"
        f"- highest_scoring_areas and lowest_scoring_areas (category names)\n"
        f"- key_strengths and key_weaknesses: 4-6 specific, detailed bullets each "
        f"(1-2 sentences per bullet, not single phrases)\n"
        f"- overall_assessment: 4-6 full paragraphs — cover the document's overall "
        f"posture, its biggest risks, how it compares to what a world-class version "
        f"would look like, and the practical consequences of its current gaps\n"
        f"- exactly 10 top_recommendations ranked by priority, each with "
        f"recommendation (detailed, 2-3 sentences), business_benefit (2-3 "
        f"sentences), risk_if_not_addressed (2-3 sentences), priority, "
        f"estimated_maturity_gain, and quick_win (true only if it takes under a "
        f"month and minimal effort)\n"
        f"- current_maturity_level and target_maturity_level (a full paragraph "
        f"describing what a Level 5 version of this specific document would contain)\n"
        f"- roadmap_to_level_5: a structured, detailed roadmap (multiple paragraphs "
        f"or phases, not a one-liner) with concrete milestones\n"
        f"- overall_readiness_pct (0-100)\n\n"
        f"Document text:\n\n{document_text}"
    )

    return system, user


def _is_complete(result: GeneratedDocumentAssessment) -> bool:
    return len(result.category_scores) >= len(DOCUMENT_MATURITY_CATEGORIES) and bool(result.top_recommendations)


def generate_document_assessment(document_title: str, document_text: str) -> GeneratedDocumentAssessment:
    system, user = document_assessment_prompt(document_title, document_text)

    result = None
    for _ in range(2):  # one retry: this is a large generation and occasionally runs long
        result = generate_structured(
            system=system, user=user, schema_model=GeneratedDocumentAssessment, max_tokens=48000, stream=True,
        )
        if _is_complete(result):
            return result

    raise ValueError(
        "The assessment response was incomplete (likely cut off generating a large "
        "document). Please try again."
    )
